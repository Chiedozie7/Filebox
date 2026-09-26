import csv
import json
import re
import sys
from collections import Counter, defaultdict
from io import StringIO
from pathlib import Path


def render_pdf(input_path, output_dir):
    import pymupdf as fitz

    output = Path(output_dir)
    pages = []
    with fitz.open(input_path) as document:
        for index, page in enumerate(document):
            image_path = output / f"page-{index + 1:04d}.png"
            page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False).save(image_path)
            pages.append(str(image_path))
    print(json.dumps(pages))


def ocr_words(tsv):
    words = []
    fields = ("level", "page_num", "block_num", "par_num", "line_num", "word_num",
              "left", "top", "width", "height", "conf", "text")
    for item in csv.DictReader(StringIO(tsv or ""), fieldnames=fields, delimiter="\t"):
        if item.get("level") != "5" or not item.get("text", "").strip():
            continue
        try:
            words.append({
                "text": item["text"].strip(),
                "x": int(item["left"]) + int(item["width"]) / 2,
                "y": int(item["top"]) + int(item["height"]) / 2,
                "line": (item["block_num"], item["par_num"], item["line_num"]),
            })
        except (KeyError, ValueError):
            continue
    return words


def detect_grids(image_path, words):
    """Accept only a repeated ruled grid with matching outer and inner lines."""
    try:
        import cv2
        import numpy as np
    except ImportError:
        return []

    image = cv2.imread(str(image_path), cv2.IMREAD_GRAYSCALE)
    if image is None:
        return []
    height, width = image.shape
    ink = cv2.adaptiveThreshold(image, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
                                cv2.THRESH_BINARY_INV, 31, 15)
    segments = cv2.HoughLinesP(ink, 1, np.pi / 1800, threshold=70,
                               minLineLength=max(50, int(width * .25)), maxLineGap=20)
    if segments is None:
        return []
    angles = [np.degrees(np.arctan2(y2 - y1, x2 - x1))
              for x1, y1, x2, y2 in segments.reshape(-1, 4)
              if abs(x2 - x1) > width * .25 and abs(y2 - y1) < abs(x2 - x1) * .1]
    if len(angles) < 3:
        return []
    matrix = cv2.getRotationMatrix2D((width / 2, height / 2), float(np.median(angles)), 1)
    straight = cv2.warpAffine(image, matrix, (width, height), borderValue=255)
    ink = cv2.adaptiveThreshold(straight, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
                                cv2.THRESH_BINARY_INV, 31, 15)

    def line_boxes(kernel, horizontal):
        mask = cv2.morphologyEx(ink, cv2.MORPH_OPEN, kernel)
        _, _, stats, _ = cv2.connectedComponentsWithStats(mask)
        boxes = []
        for x, y, line_width, line_height, _ in stats[1:]:
            if horizontal and line_width >= width * .35 and line_height <= 15:
                boxes.append((int(x), int(y), int(line_width), int(line_height)))
            elif not horizontal and line_height >= height * .08 and line_width <= 15:
                boxes.append((int(x), int(y), int(line_width), int(line_height)))
        return boxes

    horizontal = sorted(line_boxes(cv2.getStructuringElement(
        cv2.MORPH_RECT, (max(20, int(width * .09)), 1)), True), key=lambda box: box[1])
    vertical = line_boxes(cv2.getStructuringElement(
        cv2.MORPH_RECT, (1, max(20, int(height * .06)))), False)
    groups = []
    for line in horizontal:
        if (groups and height * .01 <= line[1] - groups[-1][-1][1] <= height * .08
                and abs(line[0] - groups[-1][-1][0]) <= 15
                and abs(line[2] - groups[-1][-1][2]) <= 15):
            groups[-1].append(line)
        else:
            groups.append([line])

    grids = []
    for group in groups:
        if not 3 <= len(group) <= 31:
            continue
        ys = [y + line_height / 2 for _, y, _, line_height in group]
        gaps = np.diff(ys)
        if max(gaps) > min(gaps) * 1.3:
            continue
        left = min(x for x, _, _, _ in group)
        right = max(x + line_width for x, _, line_width, _ in group)
        columns = []
        for x, y, line_width, line_height in vertical:
            center = x + line_width / 2
            if left - 10 <= center <= right + 10:
                columns.append((center, y, y + line_height))
        clusters = []
        for line in sorted(columns):
            if clusters and line[0] - clusters[-1][-1][0] <= 6:
                clusters[-1].append(line)
            else:
                clusters.append([line])
        xs = [sum(line[0] for line in cluster) / len(cluster) for cluster in clusters
              if min(line[1] for line in cluster) <= ys[0] + 8
              and max(line[2] for line in cluster) >= ys[-1] - 8]
        if not 3 <= len(xs) <= 13 or abs(xs[0] - left) > 12 or abs(xs[-1] - right) > 12:
            continue
        if right - left < width * .35 or ys[-1] - ys[0] < height * .07:
            continue
        if any(not np.any(ink[max(0, int(y) - 5):int(y) + 6,
                             max(0, int(x) - 5):int(x) + 6])
               for x in xs for y in ys):
            continue
        cells = [[[] for _ in range(len(xs) - 1)] for _ in range(len(ys) - 1)]
        table_line_counts = Counter()
        all_line_counts = Counter(word["line"] for word in words)
        for word in words:
            x = matrix[0, 0] * word["x"] + matrix[0, 1] * word["y"] + matrix[0, 2]
            y = matrix[1, 0] * word["x"] + matrix[1, 1] * word["y"] + matrix[1, 2]
            col = int(np.searchsorted(xs, x) - 1)
            row = int(np.searchsorted(ys, y) - 1)
            if 0 <= row < len(cells) and 0 <= col < len(cells[0]):
                cells[row][col].append(word)
                table_line_counts[word["line"]] += 1
        values = [[" ".join(word["text"] for word in cell)
                   for cell in row] for row in cells]
        table_lines = {line for line, count in table_line_counts.items()
                       if count == all_line_counts[line]}
        grids.append({"top": ys[0], "bottom": ys[-1], "cells": values,
                      "table_lines": table_lines, "matrix": matrix})
    return sorted(grids, key=lambda grid: grid["top"])


def add_text(document, text):
    normalized = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    for block in re.split(r"\n\s*\n", normalized):
        lines = [line.strip() for line in block.splitlines() if line.strip()]
        if lines:
            document.add_paragraph("\n".join(lines))


def write_docx(text_path, output_path):
    from docx import Document
    from docx.shared import Inches, Pt

    pages = json.loads(Path(text_path).read_text(encoding="utf-8"))
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.8)
    normal = document.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)

    for index, page in enumerate(pages):
        if index:
            document.add_page_break()
        if isinstance(page, str):
            add_text(document, page)
            continue
        text = page.get("text", "")
        words = ocr_words(page.get("tsv"))
        grids = detect_grids(page.get("imagePath"), words) if page.get("imagePath") else []
        if not grids:
            add_text(document, text)
            continue

        lines = text.replace("\r\n", "\n").replace("\r", "\n").splitlines()
        word_lines = defaultdict(list)
        for word in words:
            word_lines[word["line"]].append(word)
        cursor = 0
        for grid in grids:
            after = []
            for line, entries in word_lines.items():
                word = entries[0]
                matrix = grid["matrix"]
                y = matrix[1, 0] * word["x"] + matrix[1, 1] * word["y"] + matrix[1, 2]
                if y > grid["bottom"] + 5:
                    after.append((y, " ".join(entry["text"] for entry in entries)))
            anchor = None
            for _, label in sorted(after):
                label = re.sub(r"\s+", " ", label).casefold()
                anchor = next((position for position in range(cursor, len(lines))
                               if re.sub(r"\s+", " ", lines[position]).strip().casefold() == label), None)
                if anchor is not None:
                    break
            stop = anchor if anchor is not None else len(lines)
            table_text = {" ".join(word["text"] for word in word_lines[line]).casefold()
                          for line in grid["table_lines"]}
            add_text(document, "\n".join(line for line in lines[cursor:stop]
                                         if line.strip().casefold() not in table_text))
            table = document.add_table(rows=len(grid["cells"]), cols=len(grid["cells"][0]))
            table.style = "Table Grid"
            for row_index, row in enumerate(grid["cells"]):
                for col_index, value in enumerate(row):
                    table.cell(row_index, col_index).text = value
            cursor = stop
        add_text(document, "\n".join(lines[cursor:]))

    document.save(output_path)


if __name__ == "__main__":
    if len(sys.argv) != 4:
        raise SystemExit("Usage: ocr_to_docx.py render-pdf|write-docx input output")
    if sys.argv[1] == "render-pdf":
        render_pdf(sys.argv[2], sys.argv[3])
    elif sys.argv[1] == "write-docx":
        write_docx(sys.argv[2], sys.argv[3])
    else:
        raise SystemExit("Unsupported command")
