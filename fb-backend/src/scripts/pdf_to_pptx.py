"""Best-effort editable PDF reconstruction; no Office or OCR subprocesses."""
import json
import math
import re
import sys
from io import BytesIO

import pymupdf as fitz
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt


def rgb(color):
    if isinstance(color, int):
        return RGBColor((color >> 16) & 255, (color >> 8) & 255, color & 255)
    if len(color) == 1:
        color = color * 3
    if len(color) == 4:
        c, m, y, k = color
        color = (1 - min(1, c + k), 1 - min(1, m + k), 1 - min(1, y + k))
    return RGBColor(*(round(max(0, min(1, value)) * 255) for value in color[:3]))


def contained(rect, region, tolerance=1):
    return (rect.x0 >= region.x0 - tolerance and rect.y0 >= region.y0 - tolerance
            and rect.x1 <= region.x1 + tolerance and rect.y1 <= region.y1 + tolerance)


def opacity(drawing, key):
    return 1 if drawing.get(key) is None else drawing[key]


def seq_for(rect, log, kinds):
    matches = [(index, fitz.Rect(box)) for index, (kind, box) in enumerate(log) if kind in kinds]
    if not matches:
        return len(log)
    return max(matches, key=lambda item: (rect & item[1]).get_area())[0]


def font_for(run, span, scale):
    name = re.sub(r"^[A-Z]{6}\+", "", span.get("font", "Arial"))
    run.font.name = name
    run.font.size = Pt(max(1, min(400, span.get("size", 11) * scale)))
    flags = span.get("flags", 0)
    run.font.bold = bool(flags & 16)
    run.font.italic = bool(flags & 2)
    run.font.underline = bool(span.get("char_flags", 0) & 2)
    run.font.color.rgb = rgb(span.get("color", 0))


def reconstruct_page(prs, page, scale, offset, counts):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    blocks = page.get_text("dict")["blocks"]
    drawings = page.get_drawings()
    log = page.get_bboxlog()
    events = []
    ox, oy = offset

    def box(rect):
        return (Pt(ox + rect.x0 * scale), Pt(oy + rect.y0 * scale),
                Pt(max(.01, rect.width * scale)), Pt(max(.01, rect.height * scale)))

    def picture(content, rect, label):
        shape = slide.shapes.add_picture(BytesIO(content), *box(rect))
        shape.name = label

    # Only complete ruled grids qualify. Ambiguous/merged layouts retain text and lines.
    tables = []
    for table in page.find_tables(strategy="lines_strict").tables:
        if table.row_count < 2 or table.col_count < 2 or table.row_count * table.col_count > 2000:
            continue
        rows = table.rows
        if any(cell is None for row in rows for cell in row.cells):
            continue
        xs = [rows[0].cells[0][0]] + [cell[2] for cell in rows[0].cells]
        ys = [rows[0].cells[0][1]] + [row.cells[0][3] for row in rows]
        if any(abs(cell[0] - xs[col]) > 1 or abs(cell[2] - xs[col + 1]) > 1
               or abs(cell[1] - ys[row]) > 1 or abs(cell[3] - ys[row + 1]) > 1
               for row, cells in enumerate(rows) for col, cell in enumerate(cells.cells)):
            continue
        rect = fitz.Rect(table.bbox)
        tables.append(rect)
        events.append((seq_for(rect, log, {"fill-text", "stroke-path", "fill-path"}),
                       "table", (rect, table.extract(), xs, ys)))

    for block in blocks:
        rect = fitz.Rect(block["bbox"])
        if block["type"] == 1:
            events.append((seq_for(rect, log, {"fill-image"}), "image", block))
        elif block["type"] == 0:
            for line in block["lines"]:
                line_rect = fitz.Rect(line["bbox"])
                if any(contained(line_rect, table, 3) for table in tables):
                    continue
                events.append((seq_for(line_rect, log, {"fill-text", "stroke-text"}),
                               "text", (line, rect)))
    for drawing in drawings:
        rect = fitz.Rect(drawing["rect"])
        if any(contained(rect, table, 1) for table in tables):
            continue
        events.append((drawing.get("seqno", 0), "drawing", drawing))
    for index, (kind, bounds) in enumerate(log):
        if kind == "fill-shade":
            events.append((index, "shade", fitz.Rect(bounds)))

    vector_copy = None

    def local_vector_fallback(rect):
        nonlocal vector_copy
        # Remove text/images before clipping, so the fallback cannot duplicate
        # editable objects placed above it. Keep the original vector artwork.
        if vector_copy is None:
            vector_copy = fitz.open()
            vector_copy.insert_pdf(page.parent, from_page=page.number, to_page=page.number)
            vector_copy[0].add_redact_annot(vector_copy[0].rect, fill=False)
            vector_copy[0].apply_redactions(images=1, graphics=0, text=0)
        return vector_copy[0].get_pixmap(matrix=fitz.Matrix(2, 2), clip=rect, alpha=True).tobytes("png")

    def draw_native(drawing):
        items = drawing["items"]
        if len(items) != 1 or opacity(drawing, "fill_opacity") < 1 or opacity(drawing, "stroke_opacity") < 1:
            return False
        item = items[0]
        if item[0] == "re":
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, *box(fitz.Rect(item[1])))
        elif item[0] == "l" and not drawing.get("fill"):
            p1, p2 = item[1:3]
            shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                Pt(ox + p1.x * scale), Pt(oy + p1.y * scale),
                Pt(ox + p2.x * scale), Pt(oy + p2.y * scale))
        else:
            return False
        shape.name = "PDF editable shape"
        if item[0] == "re":
            if drawing.get("fill") is not None:
                shape.fill.solid()
                shape.fill.fore_color.rgb = rgb(drawing["fill"])
            else:
                shape.fill.background()
        if drawing.get("color") is not None:
            shape.line.color.rgb = rgb(drawing["color"])
            shape.line.width = Pt(max(.1, (drawing.get("width") or 0) * scale))
        else:
            shape.line.fill.background()
        if (drawing.get("dashes") or "[] 0").strip() not in ("[] 0", "[ ] 0"):
            from pptx.enum.dml import MSO_LINE_DASH_STYLE
            shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        return True

    def raster_path(drawing, rect):
        # Render just this path onto a transparent scratch page. This preserves
        # curves without baking nearby editable text or photos into the image.
        with fitz.open() as scratch:
            target = scratch.new_page(width=page.rect.width, height=page.rect.height)
            shape = target.new_shape()
            for item in drawing["items"]:
                if item[0] == "l":
                    shape.draw_line(*item[1:3])
                elif item[0] == "re":
                    shape.draw_rect(item[1])
                elif item[0] == "qu":
                    shape.draw_quad(item[1])
                elif item[0] == "c":
                    shape.draw_bezier(*item[1:5])
                else:
                    return local_vector_fallback(rect)
            shape.finish(color=drawing.get("color"), fill=drawing.get("fill"),
                         width=drawing.get("width") or 0, dashes=drawing.get("dashes"),
                         closePath=drawing.get("closePath", False),
                         even_odd=drawing.get("even_odd", False),
                         fill_opacity=opacity(drawing, "fill_opacity"),
                         stroke_opacity=opacity(drawing, "stroke_opacity"))
            shape.commit()
            return target.get_pixmap(matrix=fitz.Matrix(2, 2), clip=rect, alpha=True).tobytes("png")

    try:
        for _, kind, data in sorted(events, key=lambda event: event[0]):
            if kind == "text":
                line, block_rect = data
                rect = fitz.Rect(line["bbox"])
                shape = slide.shapes.add_textbox(*box(rect))
                shape.name = "PDF editable text"
                frame = shape.text_frame
                frame.clear()
                frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
                frame.word_wrap = False
                frame.auto_size = MSO_AUTO_SIZE.NONE
                paragraph = frame.paragraphs[0]
                paragraph.space_before = paragraph.space_after = Pt(0)
                paragraph.alignment = PP_ALIGN.LEFT
                if rect.width < block_rect.width * .85:
                    if abs((rect.x0 + rect.x1) - (block_rect.x0 + block_rect.x1)) < 3:
                        paragraph.alignment = PP_ALIGN.CENTER
                    elif abs(rect.x1 - block_rect.x1) < 2:
                        paragraph.alignment = PP_ALIGN.RIGHT
                for span in line["spans"]:
                    run = paragraph.add_run()
                    run.text = span["text"]
                    font_for(run, span, scale)
                    for link in page.get_links():
                        if link.get("uri") and fitz.Rect(span["bbox"]).intersects(link["from"]):
                            run.hyperlink.address = link["uri"]
                            break
                direction = line.get("dir", (1, 0))
                if abs(direction[1]) > .01:
                    shape.rotation = math.degrees(math.atan2(direction[1], direction[0])) % 360
                counts["textBoxes"] += 1
            elif kind == "image":
                content = data["image"]
                if data.get("mask"):
                    with Image.open(BytesIO(content)) as source, Image.open(BytesIO(data["mask"])) as mask:
                        rgba = source.convert("RGBA")
                        rgba.putalpha(mask.convert("L").resize(rgba.size))
                        stream = BytesIO()
                        rgba.save(stream, format="PNG")
                        content = stream.getvalue()
                # Convert formats python-pptx cannot embed while retaining a separate image.
                if data.get("ext") not in ("png", "jpeg", "jpg"):
                    with Image.open(BytesIO(content)) as source:
                        stream = BytesIO()
                        source.save(stream, format="PNG")
                        content = stream.getvalue()
                picture(content, fitz.Rect(data["bbox"]), "PDF image")
                counts["images"] += 1
            elif kind == "table":
                rect, values, xs, ys = data
                table = slide.shapes.add_table(len(values), len(xs) - 1, *box(rect)).table
                table.first_row = False
                table.horz_banding = False
                for col in range(len(xs) - 1):
                    table.columns[col].width = Pt((xs[col + 1] - xs[col]) * scale)
                for row, values_row in enumerate(values):
                    table.rows[row].height = Pt((ys[row + 1] - ys[row]) * scale)
                    for col, value in enumerate(values_row):
                        cell = table.cell(row, col)
                        cell.text = value or ""
                        cell.margin_left = cell.margin_right = Pt(2 * scale)
                        cell.margin_top = cell.margin_bottom = 0
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = RGBColor(255, 255, 255)
                        cell_rect = fitz.Rect(xs[col], ys[row], xs[col + 1], ys[row + 1])
                        spans = [span for block in blocks if block["type"] == 0
                                 for line in block["lines"] for span in line["spans"]
                                 if contained(fitz.Rect(span["bbox"]), cell_rect, 3)]
                        for drawing in drawings:
                            if drawing.get("fill") is not None and contained(cell_rect, fitz.Rect(drawing["rect"])):
                                cell.fill.fore_color.rgb = rgb(drawing["fill"])
                        for paragraph in cell.text_frame.paragraphs:
                            paragraph.space_before = paragraph.space_after = Pt(0)
                            for run in paragraph.runs:
                                font_for(run, spans[0] if spans else {"size": 10}, scale)
                        # Explicit simple borders avoid theme-dependent table styling.
                        props = cell._tc.get_or_add_tcPr()
                        for edge in ("L", "R", "T", "B"):
                            line = OxmlElement("a:ln" + edge)
                            line.set("w", "6350")
                            fill = OxmlElement("a:solidFill")
                            color = OxmlElement("a:srgbClr")
                            color.set("val", "444444")
                            fill.append(color)
                            line.append(fill)
                            props.append(line)
                counts["tables"] += 1
            elif kind == "drawing":
                if draw_native(data):
                    counts["shapes"] += 1
                else:
                    rect = (fitz.Rect(data["rect"]) + (-2, -2, 2, 2)) & page.rect
                    if not rect.is_empty:
                        picture(raster_path(data, rect), rect, "PDF localized vector fallback")
                        counts["rasterRegions"] += 1
            elif kind == "shade":
                rect = data & page.rect
                if not rect.is_empty:
                    picture(local_vector_fallback(rect), rect, "PDF shading fallback")
                    counts["rasterRegions"] += 1
    finally:
        if vector_copy is not None:
            vector_copy.close()


def convert(input_path, output_path):
    counts = dict(textBoxes=0, images=0, shapes=0, tables=0, rasterRegions=0)
    with fitz.open(input_path) as source, fitz.open() as normalized:
        if source.needs_pass:
            raise ValueError("Password-protected PDF must be unlocked first")
        if not len(source):
            raise ValueError("PDF contains no pages")
        # Normalize page rotation through a PDF form, retaining extractable objects.
        for index, page in enumerate(source):
            target = normalized.new_page(width=page.rect.width, height=page.rect.height)
            if page.get_contents():
                target.show_pdf_page(target.rect, source, index)
        widths = [page.rect.width for page in normalized]
        heights = [page.rect.height for page in normalized]
        width, height = max(widths), max(heights)
        scale = min(1, 4032 / max(width, height))
        prs = Presentation()
        prs.slide_width = Pt(max(72, width * scale))
        prs.slide_height = Pt(max(72, height * scale))
        for page in normalized:
            offset = ((prs.slide_width.pt - page.rect.width * scale) / 2,
                      (prs.slide_height.pt - page.rect.height * scale) / 2)
            reconstruct_page(prs, page, scale, offset, counts)
        prs.save(output_path)
        return {"slideCount": len(prs.slides), "reconstruction": {
            **counts, "mixedPageSizes": len(set(zip(widths, heights))) > 1}}


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: pdf_to_pptx.py input.pdf output.pptx")
    print(json.dumps(convert(sys.argv[1], sys.argv[2])))
