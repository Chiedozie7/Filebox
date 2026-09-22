import json
import sys

from docx import Document


def extract_table(table):
    column_count = len(table.columns)
    rows = []
    positions = {}
    elements = {}

    for row_index, row in enumerate(table.rows):
        cells = (
            [None] * row.grid_cols_before
            + list(row.cells)
            + [None] * row.grid_cols_after
        )
        values = []

        for column_index, cell in enumerate(cells):
            if cell is None:
                values.append("")
                continue

            key = id(cell._tc)
            if key not in positions:
                positions[key] = []
                elements[key] = cell._tc
                values.append(cell.text)
            else:
                values.append("")
            positions[key].append((row_index, column_index))

        rows.append(values + [""] * (column_count - len(values)))

    merges = []
    for occupied in positions.values():
        if len(occupied) == 1:
            continue
        first_row = min(row for row, _ in occupied)
        last_row = max(row for row, _ in occupied)
        first_col = min(col for _, col in occupied)
        last_col = max(col for _, col in occupied)
        merges.append([first_row + 1, first_col + 1, last_row + 1, last_col + 1])

    return {"rows": rows, "merges": merges}


def main():
    if len(sys.argv) != 2:
        print("Usage: docx_to_tables.py <input.docx>", file=sys.stderr)
        sys.exit(1)

    document = Document(sys.argv[1])
    tables = []
    visited = {}

    def visit(table):
        key = id(table._tbl)
        if key in visited:
            return
        visited[key] = table._tbl
        tables.append(extract_table(table))
        for row in table.rows:
            for cell in row.cells:
                for nested in cell.tables:
                    visit(nested)

    for table in document.tables:
        visit(table)

    print(json.dumps(tables))


if __name__ == "__main__":
    main()
