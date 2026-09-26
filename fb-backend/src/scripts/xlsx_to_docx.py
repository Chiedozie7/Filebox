import posixpath
import re
import sys
import zipfile
from io import BytesIO
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont


MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
OFFICE_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
DRAWING = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
DRAWING_MAIN = "http://schemas.openxmlformats.org/drawingml/2006/main"
CHART = "http://schemas.openxmlformats.org/drawingml/2006/chart"


def column_index(reference):
    letters = re.match(r"[A-Z]+", reference).group()
    result = 0
    for letter in letters:
        result = result * 26 + ord(letter) - ord("A") + 1
    return result - 1


def coordinate(reference):
    return int(re.search(r"\d+", reference).group()) - 1, column_index(reference)


def shared_strings(archive):
    if "xl/sharedStrings.xml" not in archive.namelist():
        return []
    root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
    return ["".join(node.itertext()) for node in root.findall(f"{{{MAIN}}}si")]


def cell_text(cell, strings):
    kind = cell.get("t")
    value = cell.find(f"{{{MAIN}}}v")
    formula = cell.find(f"{{{MAIN}}}f")
    if kind == "inlineStr":
        inline = cell.find(f"{{{MAIN}}}is")
        return "" if inline is None else "".join(inline.itertext())
    if value is not None and value.text is not None:
        if kind == "s":
            return strings[int(value.text)]
        if kind == "b":
            return "TRUE" if value.text == "1" else "FALSE"
        return value.text
    return f"={formula.text}" if formula is not None and formula.text else ""


def worksheet_path(target):
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join("xl", target))


def related_parts(archive, source_path):
    directory, filename = posixpath.split(source_path)
    relationships_path = posixpath.join(directory, "_rels", filename + ".rels")
    if relationships_path not in archive.namelist():
        return {}
    root = ET.fromstring(archive.read(relationships_path))
    return {
        item.get("Id"): (item.get("Type", ""),
                         item.get("Target", "").lstrip("/") if item.get("Target", "").startswith("/")
                         else posixpath.normpath(posixpath.join(directory, item.get("Target", ""))))
        for item in root.findall(f"{{{PACKAGE_REL}}}Relationship")
    }


def read_sheets(archive):
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {
        item.get("Id"): worksheet_path(item.get("Target"))
        for item in relationships.findall(f"{{{PACKAGE_REL}}}Relationship")
        if item.get("Type", "").endswith("/worksheet")
    }
    strings = shared_strings(archive)

    for sheet in workbook.findall(f"{{{MAIN}}}sheets/{{{MAIN}}}sheet"):
        relationship_id = sheet.get(f"{{{OFFICE_REL}}}id")
        root = ET.fromstring(archive.read(targets[relationship_id]))
        values = {}
        formulas = set()
        max_row = max_col = -1
        for cell in root.findall(f"{{{MAIN}}}sheetData/{{{MAIN}}}row/{{{MAIN}}}c"):
            row, col = coordinate(cell.get("r"))
            if cell.find(f"{{{MAIN}}}f") is not None:
                formulas.add((row, col))
            text = cell_text(cell, strings)
            if text:
                values[row, col] = text
                max_row, max_col = max(max_row, row), max(max_col, col)

        merges = []
        for merged in root.findall(f"{{{MAIN}}}mergeCells/{{{MAIN}}}mergeCell"):
            start, end = merged.get("ref").split(":")
            first, last = coordinate(start), coordinate(end)
            merges.append((first, last))
            max_row, max_col = max(max_row, last[0]), max(max_col, last[1])

        yield (sheet.get("name"), values, max_row + 1, max_col + 1, merges,
               formulas, targets[relationship_id], root)


def row_groups(values, merges):
    occupied = {row for row, _ in values}
    for (first_row, _), (last_row, _) in merges:
        occupied.update(range(first_row, last_row + 1))
    groups = []
    for row in sorted(occupied):
        if not groups or row > groups[-1][1]:
            groups.append([row, row + 1])
        else:
            groups[-1][1] = row + 1
    return groups


def list_items(values, formulas, merges, first_row, last_row):
    if not 2 <= last_row - first_row <= 12:
        return None
    if any(start[0] < last_row and end[0] >= first_row for start, end in merges):
        return None
    if any(first_row <= row < last_row for row, _ in formulas):
        return None
    items = []
    columns = set()
    numbers = []
    for row in range(first_row, last_row):
        cells = [(col, text) for (cell_row, col), text in values.items() if cell_row == row]
        if len(cells) != 1:
            return None
        col, text = cells[0]
        columns.add(col)
        if "\n" in text or len(text.strip()) > 80:
            return None
        match = re.fullmatch(r"\s*(?:(\d+)[.)]|[-*•])\s+(.+?)\s*", text)
        if not match:
            return None
        numbers.append(int(match.group(1)) if match.group(1) else None)
        items.append(match.group(2))
    if len(columns) != 1 or (any(number is not None for number in numbers)
                             and numbers != list(range(1, len(numbers) + 1))):
        return None
    return items


def add_table(document, values, row_count, column_count, merges, start_row=0):
    row_count = max(1, row_count)
    column_count = max(1, column_count)
    table = document.add_table(rows=row_count, cols=column_count)
    table.style = "Table Grid"
    table.autofit = False

    section = document.sections[0]
    available_width = section.page_width - section.left_margin - section.right_margin
    weights = []
    for col in range(column_count):
        longest = max(
            (len(line) for row in range(row_count)
             for line in values.get((start_row + row, col), "").splitlines()),
            default=0,
        )
        weights.append(max(8, min(24, longest)))
    total_weight = sum(weights)

    for col, weight in enumerate(weights):
        table.columns[col].width = int(available_width * weight / total_weight)
    for row in range(row_count):
        for col in range(column_count):
            cell = table.cell(row, col)
            cell.width = table.columns[col].width
            cell.text = values.get((start_row + row, col), "")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)

    for (first_row, first_col), (last_row, last_col) in merges:
        merged = table.cell(first_row - start_row, first_col).merge(
            table.cell(last_row - start_row, last_col))
        for paragraph in list(merged.paragraphs)[1:]:
            if not paragraph.text:
                paragraph._element.getparent().remove(paragraph._element)


def add_sheet(document, name, values, row_count, column_count, merges, formulas):
    document.add_heading(name, level=1)
    groups = row_groups(values, merges)
    candidates = [(first, last, list_items(values, formulas, merges, first, last))
                  for first, last in groups]
    if not any(items is not None for _, _, items in candidates):
        add_table(document, values, row_count, column_count, merges)
        return
    for first, last, items in candidates:
        if items is not None:
            for item in items:
                document.add_paragraph(item, style="List Bullet")
        else:
            block_merges = [(start, end) for start, end in merges
                            if first <= start[0] and end[0] < last]
            add_table(document, values, last - first, column_count, block_merges, first)


def chart_reference(formula, sheets):
    if not formula or "!" not in formula:
        return []
    sheet_name, addresses = formula.rsplit("!", 1)
    sheet_name = sheet_name.strip("'").replace("''", "'")
    if sheet_name not in sheets:
        return []
    endpoints = addresses.replace("$", "").split(":")
    first = coordinate(endpoints[0])
    last = coordinate(endpoints[-1])
    return [sheets[sheet_name].get((row, col), "")
            for row in range(first[0], last[0] + 1)
            for col in range(first[1], last[1] + 1)]


def chart_points(element, sheets):
    if element is None:
        return []
    formula = element.find(f".//{{{CHART}}}f")
    if formula is not None and formula.text:
        points = chart_reference(formula.text, sheets)
        if points:
            return points
    cache = element.find(f".//{{{CHART}}}strCache")
    if cache is None:
        cache = element.find(f".//{{{CHART}}}numCache")
    if cache is None:
        return []
    return [point.findtext(f"{{{CHART}}}v", "") for point in cache.findall(f"{{{CHART}}}pt")]


def chart_font(size):
    for filename in ("DejaVuSans.ttf", "arial.ttf"):
        try:
            return ImageFont.truetype(filename, size)
        except OSError:
            pass
    return ImageFont.load_default()


def chart_image(chart_xml, sheets):
    root = ET.fromstring(chart_xml)
    plot = root.find(f".//{{{CHART}}}plotArea")
    chart = next((child for child in plot if child.tag.rsplit("}", 1)[-1].endswith("Chart")
                  and child.findall(f"{{{CHART}}}ser")), None)
    if chart is None:
        raise ValueError("Excel chart has no readable series")
    kind = chart.tag.rsplit("}", 1)[-1]
    title_node = root.find(f".//{{{CHART}}}title")
    title = " ".join(node.text for node in title_node.iter()
                     if node.tag.rsplit("}", 1)[-1] == "t" and node.text) if title_node is not None else ""
    series = []
    for item in chart.findall(f"{{{CHART}}}ser"):
        name = next((node.text for node in item.findall(f".//{{{CHART}}}tx//{{{CHART}}}v") if node.text), None)
        categories = chart_points(item.find(f"{{{CHART}}}cat"), sheets)
        numbers = chart_points(item.find(f"{{{CHART}}}val"), sheets)
        if not numbers:
            numbers = chart_points(item.find(f"{{{CHART}}}yVal"), sheets)
        try:
            values = [float(value) for value in numbers]
        except ValueError:
            continue
        if values:
            series.append((name or f"Series {len(series) + 1}", categories, values))
    if not series:
        raise ValueError("Excel chart has no readable values")

    canvas = Image.new("RGB", (960, 440), "white")
    draw = ImageDraw.Draw(canvas)
    title_font, label_font = chart_font(24), chart_font(16)
    colors = ["#2878B5", "#E07835", "#4B9A69", "#8C65AC", "#C7A245"]
    if title:
        draw.text((480, 18), title, anchor="mt", fill="#222222", font=title_font)
    left, top, right, bottom = 80, 72, 925, 345
    categories = series[0][1] or [str(index + 1) for index in range(len(series[0][2]))]
    if kind == "pieChart":
        values = series[0][2]
        total = sum(max(0, value) for value in values) or 1
        angle = -90
        for index, value in enumerate(values):
            next_angle = angle + 360 * max(0, value) / total
            draw.pieslice((280, 78, 640, 438), angle, next_angle, fill=colors[index % len(colors)],
                          outline="white", width=2)
            label = str(categories[index]) if index < len(categories) else str(index + 1)
            draw.text((20, 88 + index * 27), label, fill=colors[index % len(colors)], font=label_font)
            angle = next_angle
    else:
        maximum = max(value for _, _, values in series for value in values)
        maximum = max(1, maximum) * 1.1
        for tick in range(6):
            y = bottom - (bottom - top) * tick / 5
            draw.line((left, y, right, y), fill="#DFE3E8", width=1)
            draw.text((left - 8, y), f"{maximum * tick / 5:g}", anchor="rm",
                      fill="#555555", font=label_font)
        count = max(len(values) for _, _, values in series)
        group_width = (right - left) / max(1, count)
        for index in range(count):
            label = str(categories[index]) if index < len(categories) else str(index + 1)
            draw.text((left + group_width * (index + 0.5), bottom + 10), label[:13],
                      anchor="mt", fill="#444444", font=label_font)
        for series_index, (name, _, values) in enumerate(series):
            color = colors[series_index % len(colors)]
            points = []
            for index, value in enumerate(values):
                x = left + group_width * (index + 0.5)
                y = bottom - max(0, value) / maximum * (bottom - top)
                if kind == "lineChart":
                    points.append((x, y))
                else:
                    width = group_width * 0.72 / len(series)
                    x = left + group_width * (index + 0.14) + width * series_index
                    draw.rectangle((x, y, x + width - 2, bottom), fill=color)
            if kind == "lineChart" and points:
                if len(points) > 1:
                    draw.line(points, fill=color, width=4)
                for x, y in points:
                    draw.ellipse((x - 5, y - 5, x + 5, y + 5), fill=color)
            draw.rectangle((left + series_index * 180, 405, left + series_index * 180 + 17, 420), fill=color)
            draw.text((left + series_index * 180 + 24, 405), name[:18], fill="#333333", font=label_font)
    output = BytesIO()
    canvas.save(output, format="PNG")
    output.seek(0)
    return output


def add_drawings(document, archive, sheet_path, sheet_root, sheets):
    sheet_relationships = related_parts(archive, sheet_path)
    drawings = []
    for drawing in sheet_root.findall(f"{{{MAIN}}}drawing"):
        relationship = sheet_relationships.get(drawing.get(f"{{{OFFICE_REL}}}id"))
        if not relationship or not relationship[0].endswith("/drawing"):
            continue
        drawing_path = relationship[1]
        drawing_root = ET.fromstring(archive.read(drawing_path))
        parts = related_parts(archive, drawing_path)
        for anchor in drawing_root:
            row = int(anchor.findtext(f"{{{DRAWING}}}from/{{{DRAWING}}}row", "0"))
            picture = anchor.find(f"{{{DRAWING}}}pic")
            frame = anchor.find(f"{{{DRAWING}}}graphicFrame")
            if picture is not None:
                blip = picture.find(f".//{{{DRAWING_MAIN}}}blip")
                part = parts.get(blip.get(f"{{{OFFICE_REL}}}embed")) if blip is not None else None
                if part and part[0].endswith("/image"):
                    drawings.append((row, "image", part[1]))
            elif frame is not None:
                chart = frame.find(f".//{{{CHART}}}chart")
                part = parts.get(chart.get(f"{{{OFFICE_REL}}}id")) if chart is not None else None
                if part and part[0].endswith("/chart"):
                    drawings.append((row, "chart", part[1]))
    for _, kind, part_path in sorted(drawings):
        if kind == "chart":
            document.add_picture(chart_image(archive.read(part_path), sheets), width=Inches(6.4))
        else:
            content = archive.read(part_path)
            with Image.open(BytesIO(content)) as picture:
                width = min(6.4, 2.7 * picture.width / picture.height)
            document.add_picture(BytesIO(content), width=Inches(width))


def main():
    if len(sys.argv) != 3:
        print("Usage: xlsx_to_docx.py <input.xlsx> <output.docx>", file=sys.stderr)
        sys.exit(1)

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(11.7), Inches(8.3)
    section.left_margin = section.right_margin = Inches(0.6)
    section.top_margin = section.bottom_margin = Inches(0.6)

    with zipfile.ZipFile(sys.argv[1]) as archive:
        sheets = list(read_sheets(archive))
        values_by_sheet = {sheet[0]: sheet[1] for sheet in sheets}
        for index, sheet in enumerate(sheets):
            if index:
                document.add_page_break()
            add_sheet(document, *sheet[:6])
            add_drawings(document, archive, sheet[6], sheet[7], values_by_sheet)

    document.save(sys.argv[2])
    print(len(sheets))


if __name__ == "__main__":
    main()
