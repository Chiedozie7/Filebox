import io
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from PIL import Image


BACKEND = Path(__file__).resolve().parents[1]
SOURCE = BACKEND.parent / "test-files" / "operations-workbook.xlsx"
CONVERTER = BACKEND / "src" / "scripts" / "xlsx_to_docx.py"
sys.dont_write_bytecode = True
sys.path.insert(0, str(CONVERTER.parent))
from xlsx_to_docx import MAIN, list_items, read_sheets  # noqa: E402


def convert(source, output):
    result = subprocess.run([sys.executable, str(CONVERTER), str(source), str(output)],
                            capture_output=True, text=True, check=True)
    assert result.stdout.strip() == "3"
    return Document(output)


def list_variant(source, output):
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(output, "w") as copy:
        for item in original.infolist():
            data = original.read(item.filename)
            if item.filename == "xl/worksheets/sheet3.xml":
                root = ET.fromstring(data)
                merges = root.find(f"{{{MAIN}}}mergeCells")
                for merge in list(merges):
                    if merge.get("ref") in {"A12:F12", "A13:F13", "A14:F14"}:
                        merges.remove(merge)
                merges.set("count", str(len(merges)))
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            copy.writestr(item, data)


def main():
    with tempfile.TemporaryDirectory(prefix=".fileforge-excel-word-test-", dir=BACKEND) as directory:
        directory = Path(directory)
        output = directory / "operations.docx"
        document = convert(SOURCE, output)
        with zipfile.ZipFile(SOURCE) as archive:
            sheets = list(read_sheets(archive))
            original_image = archive.read("xl/media/image.png")
        assert len(sheets) == 3 and len(document.tables) == 6
        assert [p.text for p in document.paragraphs if p.style.name == "Heading 1"] == [
            "Dispatch", "Inventory", "Summary"]
        assert len(document.inline_shapes) == 2
        assert len(document.element.xpath('.//w:br[@w:type="page"]')) == 2
        bullets = [p.text for p in document.paragraphs if p.style.name == "List Bullet"]
        expected_bullets = ["Reconcile the closing stock balance.",
                            "Review late arrivals with route leads.",
                            "Archive signed delivery receipts."]
        assert bullets == expected_bullets
        for sheet, table in zip(sheets[:2], document.tables[:2]):
            _, values, rows, columns, merges, formulas, _, _ = sheet
            assert len(table.rows) == max(1, rows)
            assert len(table.columns) == max(1, columns)
            for (row, col), text in values.items():
                assert table.cell(row, col).text == text, (sheet[0], row, col, text)
            for (first_row, first_col), (last_row, last_col) in merges:
                assert table.cell(first_row, first_col)._tc is table.cell(last_row, last_col)._tc
            for row, col in formulas:
                assert table.cell(row, col).text == values.get((row, col), "")
        _, summary, _, _, summary_merges, summary_formulas, _, _ = sheets[2]
        title, metrics, actions_heading, link = document.tables[2:]
        assert title.cell(0, 0).text == summary[0, 0]
        assert title.cell(0, 0)._tc is title.cell(0, 5)._tc
        for row in range(2, 7):
            for col in range(2):
                assert metrics.cell(row - 2, col).text == summary.get((row, col), "")
        assert actions_heading.cell(0, 0).text == summary[9, 0]
        assert actions_heading.cell(0, 0)._tc is actions_heading.cell(0, 5)._tc
        assert [link.cell(0, col).text for col in range(2)] == [
            summary.get((16, col), "") for col in range(2)]
        assert all(metrics.cell(row - 2, col).text == summary[row, col]
                   for row, col in summary_formulas if 2 <= row < 7)
        assert link.cell(0, 1).text == summary[16, 1]
        assert all(start[0] == end[0] for start, end in summary_merges)
        with zipfile.ZipFile(output) as archive:
            media = [archive.read(name) for name in archive.namelist()
                     if name.startswith("word/media/")]
        assert original_image in media, "worksheet image bytes survive in DOCX"
        chart_images = [content for content in media if content != original_image]
        assert len(chart_images) == 1
        with Image.open(io.BytesIO(chart_images[0])) as chart:
            assert chart.width >= 800 and chart.height >= 400
            assert len(chart.convert("RGB").getcolors(chart.width * chart.height) or []) > 10

        variant = directory / "unmerged-list.xlsx"
        list_variant(SOURCE, variant)
        variant_doc = convert(variant, directory / "unmerged-list.docx")
        bullets = [p.text for p in variant_doc.paragraphs if p.style.name == "List Bullet"]
        assert bullets == expected_bullets
        assert len(variant_doc.tables) > 2, "non-list ranges stay editable tables"
        assert variant_doc.tables[0].cell(2, 0).text == "Route"
        assert variant_doc.tables[1].cell(2, 0).text == "Product"
        merged_rows = [((0, 0), (0, 2)), ((1, 0), (1, 2))]
        plain_items = {(0, 0): "1. Reconcile stock", (1, 0): "2. Review arrivals"}
        assert list_items(plain_items, set(), merged_rows, 0, 2) == [
            "Reconcile stock", "Review arrivals"]
        assert list_items({**plain_items, (1, 3): "Owner"}, set(), merged_rows, 0, 2) is None
        assert list_items(plain_items, {(0, 0)}, merged_rows, 0, 2) is None
        assert list_items({(0, 0): "1. 2026-09-26", (1, 0): "2. 2026-09-27"},
                          set(), merged_rows, 0, 2) is None
        assert list_items(plain_items, set(), [((0, 0), (1, 2))], 0, 2) is None
        assert list_items({(0, 0): "Actions", (1, 0): "1. Review arrivals"},
                          set(), merged_rows, 0, 2) is None
        print("Excel-to-Word cells, merges, cached formulas, image, chart, and conservative lists passed")


if __name__ == "__main__":
    main()
