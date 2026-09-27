import json
import sys
from io import BytesIO
from pathlib import Path

import pymupdf as fitz
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Pt

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT.parent / "test-files"


def create(folder):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(720), Pt(405)
    for title in ("Dispatch overview", "Inventory follow-up"):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Pt(30), Pt(20), Pt(640), Pt(50))
        box.text = title
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    slide = prs.slides[0]
    slide.shapes.add_picture(str(FIXTURES / "landscape.jpg"), Pt(30), Pt(90), width=Pt(190))
    table = slide.shapes.add_table(3, 2, Pt(30), Pt(270), Pt(210), Pt(85)).table
    for row, values in enumerate([["Route", "Trips"], ["Marina", "40"], ["Airport", "31"]]):
        for col, value in enumerate(values):
            table.cell(row, col).text = value
    chart = CategoryChartData()
    chart.categories = ["Marina", "Airport"]
    chart.add_series("Trips", [40, 31])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Pt(270), Pt(85), Pt(410), Pt(260), chart)
    prs.save(folder / "targeted.pptx")
    with fitz.open() as pdf:
        page = pdf.new_page(width=720, height=405)
        page.insert_text((30, 40), "Editable geometry and local fallback", fontsize=18)
        page.draw_rect(fitz.Rect(30, 70, 180, 125), color=(0, .3, .8), fill=(.7, .8, 1))
        page.draw_line((220, 80), (380, 125), color=(.8, .1, .1), width=2)
        shape = page.new_shape()
        shape.draw_bezier((430, 70), (460, 160), (600, 0), (640, 130))
        shape.finish(color=(.1, .6, .2), width=3)
        shape.commit()
        page.insert_text((30, 190), "Separate editable text below the curve", fontsize=14)
        pdf.save(folder / "geometry.pdf")
    with fitz.open() as pdf:
        for title in ("Scanned page one", "Scanned page two"):
            with fitz.open() as source:
                source_page = source.new_page(width=400, height=260)
                source_page.insert_text((30, 70), title, fontsize=24)
                source_page.draw_rect(fitz.Rect(30, 100, 350, 210), color=(0, 0, 1))
                image = source_page.get_pixmap().tobytes("png")
            page = pdf.new_page(width=400, height=260)
            page.insert_image(page.rect, stream=image)
        pdf.save(folder / "scanned.pdf")


def slide_text(slide):
    parts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            parts.append(shape.text)
        if shape.has_table:
            parts.extend(cell.text for row in shape.table.rows for cell in row.cells)
    return " ".join(parts)


def verify(folder):
    source = Presentation(folder / "targeted.pptx")
    with fitz.open(folder / "office.pdf") as pdf:
        assert len(pdf) == 2
        for index, title in enumerate(("Dispatch overview", "Inventory follow-up")):
            assert title in pdf[index].get_text()
            assert abs(pdf[index].rect.width - 720) < 1 and abs(pdf[index].rect.height - 405) < 1
        assert pdf[0].get_images(), "PPTX image survives LibreOffice"
        for label in ("Marina", "Airport", "Trips", "40", "31"):
            assert label in pdf[0].get_text(), label
        assert pdf[0].get_drawings(), "chart remains rendered"
    for filename in ("mixed-content", "table-ledger"):
        deck = Presentation(folder / f"{filename}.pptx")
        with fitz.open(FIXTURES / f"{filename}.pdf") as pdf:
            assert len(deck.slides) == len(pdf)
            assert abs(deck.slide_width.pt - pdf[0].rect.width) < 1
            for index, page in enumerate(pdf):
                content = " ".join(slide_text(deck.slides[index]).split())
                for word in page.get_text().split():
                    assert word in content, (filename, index, word)
                expected = page.find_tables(strategy="lines_strict").tables
                actual = [shape.table for shape in deck.slides[index].shapes if shape.has_table]
                assert len(actual) == len(expected)
                for table, original in zip(actual, expected):
                    assert [[cell.text for cell in row.cells] for row in table.rows] == [
                        [value or "" for value in row] for row in original.extract()]
        assert any(shape.has_text_frame for slide in deck.slides for shape in slide.shapes)
    mixed = Presentation(folder / "mixed-content.pptx")
    assert any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for slide in mixed.slides for shape in slide.shapes)
    geometry = Presentation(folder / "geometry.pptx")
    shapes = geometry.slides[0].shapes
    assert sum(s.shape_type in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.LINE) for s in shapes) >= 2
    fallback = [s for s in shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert fallback and all(s.width < geometry.slide_width / 2 for s in fallback)
    assert "Separate editable text below the curve" in slide_text(geometry.slides[0])
    scanned = Presentation(folder / "scanned.pptx")
    assert len(scanned.slides) == 2
    images = [s.image.blob for slide in scanned.slides for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(images) == 2 and images[0] != images[1]
    print(json.dumps({"verified": "editable text, exact table cells, images, basic shapes, local fallback, slide order, scanned pages, Office export"}))


if __name__ == "__main__":
    folder = Path(sys.argv[2])
    (create if sys.argv[1] == "create" else verify)(folder)
