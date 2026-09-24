"""Independent content checks of outputs produced by regression.test.js."""
import json
import re
from pathlib import Path
from zipfile import ZipFile
from lxml import etree as ET
import pymupdf
import pdfplumber
from docx import Document

root = Path(__file__).resolve().parents[2]
out = root / 'test-output/regression-20260924'
fixtures = root / 'test-files'
ns = {'x': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
      'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
results = []

def check(name, fn):
    try:
        detail = fn()
        results.append(dict(name=name, passed=True, detail=detail))
    except Exception as exc:
        results.append(dict(name=name, passed=False, error=str(exc)))

def read_xlsx(file):
    sheets = []
    with ZipFile(file) as z:
        assert z.testzip() is None
        workbook = ET.fromstring(z.read('xl/workbook.xml'))
        strings = []
        if 'xl/sharedStrings.xml' in z.namelist():
            strings = [''.join(n.xpath('.//x:t/text()', namespaces=ns)) for n in ET.fromstring(z.read('xl/sharedStrings.xml')).findall('x:si', ns)]
        for i, sheet in enumerate(workbook.findall('x:sheets/x:sheet', ns), 1):
            xml = ET.fromstring(z.read(f'xl/worksheets/sheet{i}.xml'))
            cells = {}
            for c in xml.findall('.//x:sheetData/x:row/x:c', ns):
                value = c.find('x:v', ns)
                formula = c.find('x:f', ns)
                text = value.text if value is not None else None
                if c.get('t') == 's': text = strings[int(text)]
                elif c.get('t') == 'inlineStr': text = ''.join(c.xpath('.//x:t/text()', namespaces=ns))
                elif text is None and formula is not None: text = '=' + formula.text
                cells[c.get('r')] = text or ''
            sheets.append((sheet.get('name'), cells, xml))
    return sheets

def table_checks():
    with pdfplumber.open(fixtures / 'mixed-content.pdf') as pdf:
        expected = [table for page in pdf.pages for table in page.extract_tables()]
    actual = read_xlsx(out / 'pdf-to-excel.xlsx')
    assert len(actual) == len(expected)
    count = 0
    for (_, cells, xml), table in zip(actual, expected):
        for r, row in enumerate(table, 1):
            for c, text in enumerate(row):
                assert cells.get(f'{chr(65+c)}{r}', '') == (text or '')
                count += 1
        assert all(10 <= float(c.get('width')) <= 40 for c in xml.findall('x:cols/x:col', ns))
    tables = Document(fixtures / 'operations-report.docx').tables
    actual = read_xlsx(out / 'word-to-excel.xlsx')
    assert len(actual) == len(tables)
    for (_, cells, _), table in zip(actual, tables):
        for r, row in enumerate(table.rows, 1):
            for c, cell in enumerate(row.cells):
                assert cells.get(f'{chr(65+c)}{r}', '') == cell.text
                count += 1
    return dict(exact_cells=count, pdf_tables=len(expected), docx_tables=len(tables))

def excel_word():
    expected = read_xlsx(fixtures / 'operations-workbook.xlsx')
    document = Document(out / 'excel-to-word.docx')
    assert len(document.tables) == len(expected) == 3
    headings = '\n'.join(p.text for p in document.paragraphs)
    count = 0
    for (name, cells, _), table in zip(expected, document.tables):
        assert name in headings
        for ref, value in cells.items():
            # A merged Word cell exposes its anchor text at every covered coordinate.
            if not value: continue
            letters, row = re.match(r'([A-Z]+)(\d+)', ref).groups()
            col = 0
            for char in letters: col = col * 26 + ord(char) - 64
            assert table.cell(int(row)-1, col-1).text == value, (name, ref, value)
            count += 1
    return dict(worksheets=3, exact_cells=count)

def word_quality():
    file = out / 'pdf-to-word.docx'
    document = Document(file)
    assert len(document.inline_shapes) >= 1 and len(document.tables) >= 1
    with ZipFile(file) as z:
        xml = ET.fromstring(z.read('word/document.xml'))
        links = xml.findall('.//w:p/w:hyperlink', ns)
        assert links and not xml.findall('.//w:r/w:hyperlink', ns)
        relationships = ET.fromstring(z.read('word/_rels/document.xml.rels'))
        urls = [r.get('Target') for r in relationships if r.get('Type').endswith('/hyperlink')]
        source_urls = [l['uri'] for p in pymupdf.open(fixtures/'mixed-content.pdf') for l in p.get_links() if 'uri' in l]
        assert all(url in urls for url in source_urls)
        assert all(''.join(link.xpath('.//w:t/text()', namespaces=ns)).strip() for link in links)
    return dict(images=len(document.inline_shapes), tables=len(document.tables), hyperlinks=len(links))

def pdf_quality():
    src = pymupdf.open(fixtures/'mixed-content.pdf')
    for name in ['pdf-compress.pdf', 'pdf-unlock.pdf', 'pdf-already-open.pdf']:
        doc = pymupdf.open(out/name)
        assert not doc.needs_pass and len(doc) == len(src)
        for a, b in zip(src, doc):
            assert a.get_text() == b.get_text()
            assert len(a.get_images()) == len(b.get_images())
    word = pymupdf.open(out/'word-to-pdf.pdf')
    excel = pymupdf.open(out/'excel-to-pdf.pdf')
    assert len(word) == 3 and len(excel) == 3
    assert all(p.get_text().strip() for p in excel)
    text = ''.join(p.get_text() for p in excel)
    for label in ['Dispatch', 'Inventory', 'Capacity']:
        assert label in text
    assert any(p.get_links() for p in excel) and any(p.get_images() for p in excel)
    merged = pymupdf.open(out/'mixed-merge.pdf')
    assert len(merged) == len(word) + 1 + len(src) + len(excel)
    for i, page in enumerate(word): assert merged[i].get_text() == page.get_text()
    assert merged[len(word)].get_images()
    for i, page in enumerate(src): assert merged[len(word)+1+i].get_text() == page.get_text()
    for i, page in enumerate(excel): assert merged[len(word)+1+len(src)+i].get_text() == page.get_text()
    return dict(word_pages=len(word), excel_pages=len(excel), merged_pages=len(merged), order='DOCX -> image -> PDF -> XLSX')

def ocr_quality():
    img = Document(out/'ocr-delivery-note.png.docx')
    text = '\n'.join(p.text for p in img.paragraphs)
    for phrase in ['FILEFORGE DELIVERY NOTE', 'LAGOS 2048', 'Solar panel kits', 'Marina warehouse']:
        assert phrase.lower() in text.lower(), phrase
    pdf = Document(out/'ocr-mixed-content.pdf.docx')
    xml = pdf._element
    assert len(xml.xpath('.//w:br[@w:type="page"]')) == 2
    content = '\n'.join(p.text for p in pdf.paragraphs)
    assert len(content) > 500
    return dict(image_text_characters=len(text), pdf_text_characters=len(content), pdf_page_breaks=2)

check('PDF/Word tables exact cell preservation', table_checks)
check('Excel-to-Word exact cells and worksheet names', excel_word)
check('PDF-to-Word images tables and clickable relationship', word_quality)
check('PDF contents pagination and mixed merge order', pdf_quality)
check('OCR text and PDF page separation', ocr_quality)
(out/'artifact-results.json').write_text(json.dumps(results, indent=2))
print(json.dumps(results, indent=2))
raise SystemExit(any(not item['passed'] for item in results))
