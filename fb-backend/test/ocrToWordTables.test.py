import json
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from PIL import Image


BACKEND = Path(__file__).resolve().parents[1]
FIXTURES = BACKEND.parent / "test-files" / "manual-QA"
sys.dont_write_bytecode = True
sys.path.insert(0, str(BACKEND / "src" / "scripts"))
from ocr_to_docx import detect_grids, write_docx  # noqa: E402
NODE = """
const ocr = require('./src/services/ocrService');
(async () => {
  await ocr.convertToWord(process.argv[1], process.argv[2], 'image');
  await ocr.convertToWord(process.argv[3], process.argv[4], 'pdf');
})().catch(error => { console.error(error); process.exitCode = 1; });
"""


def text(document):
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def body_items(document):
    items = []
    for element in document.element.body:
        if element.tag.endswith("}tbl"):
            items.append("TABLE")
        elif element.tag.endswith("}p"):
            items.append("".join(node.text or "" for node in element.xpath(".//w:t")))
    return items


def main():
    with tempfile.TemporaryDirectory(prefix=".fileforge-ocr-table-test-", dir=BACKEND) as directory:
        image_output = Path(directory) / "image.docx"
        pdf_output = Path(directory) / "pdf.docx"
        subprocess.run([
            "node", "-e", NODE,
            str(FIXTURES / "skewed-lowres-delivery.png"), str(image_output),
            str(FIXTURES / "image-heavy-scanned-3-pages.pdf"), str(pdf_output),
        ], cwd=BACKEND, check=True, timeout=240)

        image = Document(image_output)
        assert len(image.tables) == 1
        assert len(image.tables[0].rows) == 4 and len(image.tables[0].columns) == 3
        assert all(not cell.text for row in image.tables[0].rows for cell in row.cells)
        image_text = text(image)
        for phrase in ("DISPATCH EXCEPTION REPORT", "Route: Marina", "Scheduled trips: 40",
                       "Action: Reassign spare van", "signed delivery register"):
            assert phrase in image_text, phrase
        image_items = body_items(image)
        assert image_items.index("TABLE") < next(i for i, item in enumerate(image_items)
                                                   if "Page 2" in item)
        grid = detect_grids(FIXTURES / "skewed-lowres-delivery.png", [
            {"text": "SKU", "x": 150, "y": 690, "line": ("1", "1", "1")},
            {"text": "42", "x": 380, "y": 740, "line": ("1", "1", "2")},
        ])[0]
        assert grid["cells"][0] == ["SKU", "", ""]
        assert grid["cells"][1] == ["", "42", ""]

        pdf = Document(pdf_output)
        assert len(pdf.tables) == 3
        assert all(len(table.rows) == 4 and len(table.columns) == 3 for table in pdf.tables)
        assert all(not cell.text for table in pdf.tables for row in table.rows for cell in row.cells)
        assert len(pdf.element.xpath('.//w:br[@w:type="page"]')) == 2
        pdf_text = text(pdf)
        headings = ("WAREHOUSE INTAKE SHEET", "DISPATCH EXCEPTION REPORT", "SUPERVISOR SIGN-OFF")
        positions = [pdf_text.index(heading) for heading in headings]
        assert positions == sorted(positions)
        for phrase in ("FileForge Logistics", "Solar panel kits", "Route: Marina",
                       "Completed on time: 36", "Closing balance: 24 kits", "Approved by: Ada Okafor"):
            assert phrase in pdf_text, phrase
        items = body_items(pdf)
        table_positions = [i for i, item in enumerate(items) if item == "TABLE"]
        assert len(table_positions) == 3
        dispatch = next(i for i, item in enumerate(items) if "DISPATCH EXCEPTION REPORT" in item)
        signoff = next(i for i, item in enumerate(items) if "SUPERVISOR SIGN-OFF" in item)
        assert table_positions[0] < dispatch < table_positions[1] < signoff < table_positions[2]

        blank = Path(directory) / "blank.png"
        Image.new("RGB", (720, 950), "white").save(blank)
        plain_pages = Path(directory) / "plain.json"
        plain_pages.write_text(json.dumps([{"text": "Only text\nAnother line",
                                           "tsv": "", "imagePath": str(blank)}]), encoding="utf-8")
        plain_output = Path(directory) / "plain.docx"
        write_docx(plain_pages, plain_output)
        plain = Document(plain_output)
        assert len(plain.tables) == 0 and text(plain) == "Only text\nAnother line"
        print("OCR-to-Word image/PDF text, editable 4x3 grids, empty cells, order, and page breaks passed")


if __name__ == "__main__":
    main()
